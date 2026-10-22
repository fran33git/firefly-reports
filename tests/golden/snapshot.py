"""Snapshot helpers for the golden characterization tests.

``to_jsonable`` turns report data into tagged JSON, ``story_recorder`` records what
ReportLab is asked to draw, ``xlsx_snapshot`` reads a workbook back, ``assert_golden``
compares against (or, with ``FIREFLY_UPDATE_GOLDEN=1``, rewrites) the committed files.
"""

from __future__ import annotations

import hashlib
import json
import os
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

GOLDEN_ROOT = Path(__file__).parent
UPDATE_ENV = "FIREFLY_UPDATE_GOLDEN"


def to_jsonable(obj: object) -> object:
    """Tagged, order-preserving JSON form: Decimal/float/date keep their type visible."""
    if isinstance(obj, bool) or obj is None or isinstance(obj, (int, str)):
        return obj
    if isinstance(obj, Decimal):
        return f"D:{obj}"
    if isinstance(obj, float):
        return f"F:{obj!r}"
    if isinstance(obj, datetime):
        return f"dt:{obj.isoformat()}"
    if isinstance(obj, date):
        return f"date:{obj.isoformat()}"
    if isinstance(obj, dict):
        return {str(k): to_jsonable(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [to_jsonable(v) for v in obj]
    if isinstance(obj, (set, frozenset)):
        return sorted((to_jsonable(v) for v in obj), key=repr)
    return f"X:{type(obj).__name__}:{obj!r}"


# ---- PDF story recording ---------------------------------------------------------------


@dataclass
class PdfSnapshot:
    stem: str
    story: list[Any] = field(default_factory=list)
    strings: list[str] = field(default_factory=list)
    pages: int = 0

    def as_dict(self) -> dict[str, Any]:
        return {"story": self.story, "strings": self.strings, "pages": self.pages}


def _num(value: object) -> object:
    return round(float(value), 2) if isinstance(value, (int, float)) else value  # type: ignore[arg-type]


def _cell(cell: object) -> object:
    from reportlab.platypus import Paragraph

    if isinstance(cell, Paragraph):
        return cell.getPlainText()
    if cell is None or isinstance(cell, (str, int, float)):
        return cell
    return describe(cell)


def describe(flowable: object) -> object:
    """Plain, version-robust description of one flowable."""
    from reportlab.platypus import Image, KeepTogether, Paragraph, Spacer, Table

    name = type(flowable).__name__
    if isinstance(flowable, Paragraph):
        return {"p": flowable.style.name, "text": flowable.getPlainText()}
    if isinstance(flowable, Table):
        return {
            "table": [[_cell(c) for c in row] for row in flowable._cellvalues],
            "colWidths": [_num(w) for w in (flowable._colWidths or [])],
        }
    if isinstance(flowable, Spacer):
        return {"spacer": _num(flowable.height)}
    if isinstance(flowable, Image):
        return {"image": [_num(flowable.drawWidth), _num(flowable.drawHeight)]}
    if isinstance(flowable, KeepTogether):
        return {"keep": [describe(f) for f in flowable._content]}
    return {"flowable": name}


@contextmanager
def story_recorder() -> Iterator[list[PdfSnapshot]]:
    """Record every document built inside the block (class-level patch, survives code moves)."""
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.platypus.doctemplate import BaseDocTemplate

    snapshots: list[PdfSnapshot] = []
    state: dict[str, PdfSnapshot | None] = {"current": None}
    orig_build = BaseDocTemplate.build
    originals = {
        n: getattr(Canvas, n) for n in ("drawString", "drawRightString", "drawCentredString")
    }

    def build(self: Any, flowables: Any, *args: Any, **kwargs: Any) -> Any:
        snap = PdfSnapshot(stem=Path(str(self.filename)).stem)
        snap.story = [describe(f) for f in flowables]
        state["current"] = snap
        snapshots.append(snap)
        try:
            return orig_build(self, flowables, *args, **kwargs)
        finally:
            snap.pages = int(self.page)
            state["current"] = None

    def wrap(name: str) -> Any:
        original = originals[name]

        def draw(self: Any, x: float, y: float, text: str, *a: Any, **k: Any) -> Any:
            current = state["current"]
            if current is not None:
                current.strings.append(str(text))
            return original(self, x, y, text, *a, **k)

        return draw

    BaseDocTemplate.build = build  # type: ignore[method-assign]
    for name in originals:
        setattr(Canvas, name, wrap(name))
    try:
        yield snapshots
    finally:
        BaseDocTemplate.build = orig_build  # type: ignore[method-assign]
        for name, original in originals.items():
            setattr(Canvas, name, original)


# ---- Excel -------------------------------------------------------------------------------


def xlsx_snapshot(path: Path) -> dict[str, Any]:
    """Per sheet: layout facts and, for every non-empty cell, value and visible styling."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    sheets: dict[str, Any] = {}
    for ws in wb.worksheets:
        cells = []
        for row in ws.iter_rows():
            for c in row:
                if c.value is None and c.fill.fgColor.rgb in (None, "00000000"):
                    continue
                cells.append(
                    [
                        c.coordinate,
                        to_jsonable(c.value),
                        c.number_format,
                        bool(c.font.bold),
                        bool(c.font.italic),
                        c.font.size,
                        str(c.fill.fgColor.rgb),
                        c.alignment.horizontal,
                        c.alignment.vertical,
                        c.border.bottom.style if c.border.bottom else None,
                    ]
                )
        sheets[ws.title] = {
            "merged": sorted(str(r) for r in ws.merged_cells.ranges),
            "freeze": ws.freeze_panes,
            "filter": ws.auto_filter.ref,
            "widths": {k: _num(v.width) for k, v in sorted(ws.column_dimensions.items())},
            "cells": cells,
        }
    return {"sheets": list(sheets), "data": sheets}


# ---- Golden files ------------------------------------------------------------------------


def _first_diff(a: object, b: object, where: str = "$") -> str:
    if type(a) is not type(b):
        return f"{where}: {a!r} != {b!r}"
    if isinstance(a, dict) and isinstance(b, dict):
        for k in a.keys() | b.keys():
            if k not in a or k not in b:
                return f"{where}.{k}: key present on one side only"
            d = _first_diff(a[k], b[k], f"{where}.{k}")
            if d:
                return d
        return "" if list(a) == list(b) else f"{where}: key order differs"
    if isinstance(a, list) and isinstance(b, list):
        if len(a) != len(b):
            return f"{where}: length {len(a)} != {len(b)}"
        for i, (x, y) in enumerate(zip(a, b, strict=True)):
            d = _first_diff(x, y, f"{where}[{i}]")
            if d:
                return d
        return ""
    return "" if a == b else f"{where}: {a!r} != {b!r}"


def _dump(payload: object) -> str:
    return json.dumps(payload, ensure_ascii=False, indent=1) + "\n"


def assert_golden(name: str, payload: object, *, digest_only: bool = False) -> None:
    """Compare *payload* with ``tests/golden/<name>.json`` (or a ``.sha256`` digest)."""
    text = _dump(payload)
    update = os.environ.get(UPDATE_ENV) == "1"
    if digest_only:
        path = GOLDEN_ROOT / f"{name}.sha256"
        digest = hashlib.sha256(text.encode("utf-8")).hexdigest() + "\n"
        if update:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(digest, encoding="utf-8")
            return
        if not path.exists():
            raise AssertionError(f"golden file missing: {path} (set {UPDATE_ENV}=1)")
        if path.read_text(encoding="utf-8") != digest:
            raise AssertionError(f"golden digest differs: {name}")
        return
    path = GOLDEN_ROOT / f"{name}.json"
    if update:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")
        return
    if not path.exists():
        raise AssertionError(f"golden file missing: {path} (set {UPDATE_ENV}=1)")
    expected = json.loads(path.read_text(encoding="utf-8"))
    actual = json.loads(text)
    diff = _first_diff(expected, actual)
    if diff:
        raise AssertionError(f"golden differs: {name}: {diff}")
