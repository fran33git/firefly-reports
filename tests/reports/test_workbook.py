"""write_workbook writes sheets in the given order and skips empty data."""

from openpyxl import load_workbook

from firefly_reports.reports.workbook import write_workbook


def _writer(title):
    def write(wb, data):
        wb.create_sheet(title)

    return write


def test_write_workbook_skips_empty_data(tmp_path):
    path = tmp_path / "w.xlsx"
    write_workbook([(_writer("A"), {"x": 1}), (_writer("B"), None)], path)
    assert load_workbook(path).sheetnames == ["A"]


def test_write_workbook_keeps_the_given_order(tmp_path):
    path = tmp_path / "w.xlsx"
    write_workbook([(_writer("Z"), [1]), (_writer("A"), {"k": 1}), (_writer("M"), {})], path)
    assert load_workbook(path).sheetnames == ["Z", "A"]
