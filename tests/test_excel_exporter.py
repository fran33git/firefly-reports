"""Smoke test for render_all_xlsx_full."""

from firefly_reports.data_processor import (
    build_account_statements,
    build_bills_report,
    build_budget_vs_actual,
    build_cash_flow,
    build_cumulative_cashflow,
    build_expense_trend,
    build_income_concentration,
    build_income_expense,
    build_kpi_scorecard,
    build_liabilities_report,
    build_net_worth,
    build_savings_goals,
    build_summary,
    build_tagged_report,
    build_tax_summary,
    build_transaction_register,
    build_yoy_comparison,
)
from firefly_reports.excel_exporter import render_all_xlsx_full

MIN_XLSX_BYTES = 5000


def test_render_all_xlsx_full(
    tmp_path,
    txn_2025,
    txn_2024,
    accounts,
    liabilities,
    budgets,
    budget_limits,
    bills,
    bill_txn,
    piggy_banks,
    about,
    about_user,
    period,
    period_prev,
    owner,
    currency,
):
    start, end = period
    start_prev, end_prev = period_prev
    txn_by_account = {"1": txn_2025}

    data = {
        "cf": build_cash_flow(txn_2025, start, end, owner, currency),
        "ie": build_income_expense(txn_2025, start, end, owner, currency),
        "reg": build_transaction_register(txn_2025, start, end, owner, currency),
        "net_worth": build_net_worth(accounts, end, owner, currency),
        "statements": build_account_statements(
            accounts[:1], txn_by_account, start, end, owner, currency
        ),
        "tax": build_tax_summary(txn_2025, 2025, owner, currency),
        "trend": build_expense_trend(txn_2025, start, end, owner, currency),
        "tagged": build_tagged_report(txn_2025, ["tax"], start, end, owner, currency),
        "budget": build_budget_vs_actual(
            budgets, budget_limits, txn_2025, start, end, owner, currency
        ),
        "bills": build_bills_report(bills, bill_txn, start, end, owner, currency),
        "savings": build_savings_goals(piggy_banks, end, owner, currency),
        "liabilities": build_liabilities_report(liabilities, {}, end, start, end, owner, currency),
        "kpi": build_kpi_scorecard(txn_2025, accounts, liabilities, start, end, owner, currency),
        "yoy": build_yoy_comparison(
            txn_2025, txn_2024, start, end, start_prev, end_prev, owner, currency
        ),
        "cumulative_cf": build_cumulative_cashflow(txn_2025, start, end, owner, currency),
        "concentration": build_income_concentration(txn_2025, start, end, owner, currency),
        "summary": build_summary(
            txn_2025,
            accounts,
            liabilities,
            budgets,
            bills,
            piggy_banks,
            about,
            about_user,
            start,
            end,
            owner,
            currency,
        ),
    }

    p = tmp_path / "reports.xlsx"
    render_all_xlsx_full(data, str(p))

    assert p.exists(), "Excel file not created"
    assert p.stat().st_size > MIN_XLSX_BYTES, f"Excel too small: {p.stat().st_size} bytes"


def test_render_all_xlsx_full_skips_none_sheets(
    tmp_path,
    txn_2025,
    accounts,
    liabilities,
    period,
    owner,
    currency,
):
    """Range mode: yoy (and any other unavailable report) is passed as None."""
    start, end = period

    data = {
        "cf": build_cash_flow(txn_2025, start, end, owner, currency),
        "ie": build_income_expense(txn_2025, start, end, owner, currency),
        "yoy": None,  # year-only report, unavailable in range mode
        "concentration": build_income_concentration(txn_2025, start, end, owner, currency),
    }

    p = tmp_path / "reports_range.xlsx"
    render_all_xlsx_full(data, str(p))

    assert p.exists(), "Excel file not created"

    from openpyxl import load_workbook

    wb = load_workbook(p)
    assert len(wb.sheetnames) == 3  # cf, ie, concentration — yoy skipped


def test_render_all_xlsx_full_summary_sheet(
    tmp_path,
    txn_2025,
    accounts,
    liabilities,
    budgets,
    bills,
    piggy_banks,
    about,
    about_user,
    period,
    owner,
    currency,
):
    start, end = period
    data = {
        "summary": build_summary(
            txn_2025,
            accounts,
            liabilities,
            budgets,
            bills,
            piggy_banks,
            about,
            about_user,
            start,
            end,
            owner,
            currency,
        ),
    }
    p = tmp_path / "reports_summary.xlsx"
    render_all_xlsx_full(data, str(p))
    assert p.exists(), "Excel file not created"

    from openpyxl import load_workbook

    wb = load_workbook(p)
    assert "Summary" in wb.sheetnames

    ws = wb["Summary"]
    counts = data["summary"]["counts"]
    label_col = {row[0].value: row for row in ws.iter_rows(min_col=1, max_col=2)}
    row = label_col["Asset accounts"]
    assert row[1].value == str(counts["asset_accounts"])
    assert row[1].number_format != "#,##0.00"
