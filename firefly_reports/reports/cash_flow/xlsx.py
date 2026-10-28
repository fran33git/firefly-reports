"""Excel sheet of the cash_flow report."""

from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from firefly_reports import i18n
from firefly_reports.i18n import t
from firefly_reports.render.fmt import period_str
from firefly_reports.render.xlsx.layout import (
    data_row,
    header_row,
    title_block,
    to_number,
)
from firefly_reports.render.xlsx.styles import (
    LIGHT_HEX,
)

# ─────────────────────────────────────────────────
# Monochrome palette
# ─────────────────────────────────────────────────


# ─────────────────────────────────────────────────
# Sheet 1 — Cash Flow Statement
# ─────────────────────────────────────────────────


def write_sheet(wb: Workbook, data: dict[str, Any]) -> None:
    ws = wb.create_sheet(t("excel.sheet_cashflow"))
    period = period_str(data["period_start"], data["period_end"])
    sym = data["currency"]

    row = title_block(ws, t("excel.sheet_cashflow"), data["owner"], period)

    # ── Income
    ws.merge_cells(f"A{row}:H{row}")
    c = ws.cell(row=row, column=1, value=t("excel.section_inflows"))
    c.font = Font(name="Calibri", bold=True, size=9, color=LIGHT_HEX)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 22
    row += 1

    row = header_row(
        ws, row, [t("common.category_item"), f"{t('common.amount')} ({sym})"], col_widths=[42, 18]
    )
    for item in data["inflows"]:
        row = data_row(ws, row, [item["category"], to_number(item["amount"])], indent_col=1)
    row = data_row(ws, row, [t("common.total_income"), to_number(data["total_in"])], is_total=True)

    row += 1  # spacer row

    # ── Expenses
    ws.merge_cells(f"A{row}:H{row}")
    c = ws.cell(row=row, column=1, value=t("excel.section_outflows"))
    c.font = Font(name="Calibri", bold=True, size=9, color=LIGHT_HEX)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 22
    row += 1

    row = header_row(ws, row, [t("common.category_item"), f"{t('common.amount')} ({sym})"])
    for item in data["outflows"]:
        row = data_row(ws, row, [item["category"], to_number(item["amount"])], indent_col=1)
    row = data_row(
        ws, row, [t("common.total_expense"), to_number(data["total_out"])], is_total=True
    )

    row += 1

    # ── Monthly
    if len(data["by_month"]) > 1:
        months_it = i18n.current().section("months")
        ws.merge_cells(f"A{row}:H{row}")
        c = ws.cell(row=row, column=1, value=t("excel.section_monthly"))
        c.font = Font(name="Calibri", bold=True, size=9, color=LIGHT_HEX)
        ws.row_dimensions[row].height = 22
        row += 1

        row = header_row(
            ws,
            row,
            [
                t("common.month"),
                f"{t('common.income')} ({sym})",
                f"{t('common.expense')} ({sym})",
                f"{t('common.net_flow')} ({sym})",
            ],
            col_widths=[22, 18, 18, 20],
        )
        for mk, vals in data["by_month"].items():
            y, m = mk.split("-")
            row = data_row(
                ws,
                row,
                [
                    f"{months_it.get(m, m)} {y}",
                    to_number(vals["in"]),
                    to_number(vals["out"]),
                    to_number(vals["net"]),
                ],
            )
        row += 1

    # ── Grand total row
    row = data_row(ws, row, [t("common.net_flow"), to_number(data["net"])], is_grand=True)

    ws.freeze_panes = "A6"
    ws.sheet_view.showGridLines = False
