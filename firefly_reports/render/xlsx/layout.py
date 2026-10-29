"""Sheet layout helpers: title block, header and data rows, sheet names."""

from collections.abc import Sequence
from decimal import Decimal
from typing import Any

from openpyxl.styles import Alignment, Border, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from firefly_reports import clock
from firefly_reports.i18n import t
from firefly_reports.render.fmt import fmt_date
from firefly_reports.render.xlsx.styles import (
    DARK_HEX,
    LIGHT_HEX,
    MID_HEX,
    RULE_HEX,
    body_font,
    grand_border,
    medium,
    thin,
    total_fill,
)


def title_block(ws: Worksheet, sheet_title: str, owner: str, period: str) -> int:
    """
    Write the first 4 header rows.
    Returns the row number from which content should start.
    """
    ws.row_dimensions[1].height = 28
    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = sheet_title.upper()
    c.font = Font(name="Calibri", bold=True, size=14, color=DARK_HEX)
    c.alignment = Alignment(horizontal="left", vertical="center")
    c.border = Border(bottom=thin())

    ws.row_dimensions[2].height = 16
    ws.merge_cells("A2:H2")
    c2 = ws["A2"]
    c2.value = owner
    c2.font = Font(name="Calibri", size=10, color=MID_HEX)
    c2.alignment = Alignment(horizontal="left", vertical="center")

    ws.row_dimensions[3].height = 14
    ws.merge_cells("A3:H3")
    c3 = ws["A3"]
    c3.value = t("excel.header_period_ref").format(period=period, date=fmt_date(clock.now()))
    c3.font = Font(name="Calibri", size=9, color=LIGHT_HEX, italic=True)
    c3.alignment = Alignment(horizontal="left", vertical="center")
    c3.border = Border(bottom=medium())

    ws.row_dimensions[4].height = 8  # spacer row

    return 5  # first content row


def header_row(
    ws: Worksheet, row: int, labels: list[Any], col_widths: Sequence[float] | None = None
) -> int:
    """Write a column header row with bold font and thin bottom border."""
    for col_idx, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=col_idx, value=label)
        c.font = body_font(bold=True, size=9.5)
        c.alignment = Alignment(horizontal="right" if col_idx > 1 else "left", vertical="center")
        c.border = Border(bottom=thin(DARK_HEX))
    if col_widths:
        for i, w in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 20
    return row + 1


def data_row(
    ws: Worksheet,
    row: int,
    values: list[Any],
    indent_col: int | None = None,
    is_total: bool = False,
    is_grand: bool = False,
) -> int:
    """
    Write a data row.
    - is_total: bold + grey background + thin top border
    - is_grand: larger bold + thick top/bottom border
    """
    font_size = 10.5 if is_grand else 10
    bold = is_total or is_grand

    for col_idx, val in enumerate(values, start=1):
        c = ws.cell(row=row, column=col_idx, value=val)
        c.font = body_font(bold=bold, size=font_size)
        is_num = isinstance(val, (int, float)) and col_idx > 1
        c.alignment = Alignment(
            horizontal="right" if (is_num or col_idx > 1) else "left",
            vertical="center",
            indent=1
            if (indent_col and col_idx == indent_col and not is_total and not is_grand)
            else 0,
        )
        if is_num:
            c.number_format = "#,##0.00"
        if is_total:
            c.fill = total_fill()
            if col_idx == 1:
                c.border = Border(top=thin(RULE_HEX), bottom=thin(RULE_HEX))
            else:
                c.border = Border(top=thin(RULE_HEX), bottom=thin(RULE_HEX))
        if is_grand:
            c.fill = total_fill()
            c.border = grand_border()

    ws.row_dimensions[row].height = 18 if not is_grand else 22
    return row + 1


def safe_sheet_name(name: str, max_len: int = 31) -> str:
    """Remove characters forbidden in Excel sheet names: [ ] * ? / \\ :"""
    for ch in r"[]*?/\\:":
        name = name.replace(ch, "")
    return name[:max_len].strip() or "Sheet"


def to_number(v: Decimal) -> float:
    return float(v)
