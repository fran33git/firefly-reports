"""Colours, borders and fonts of the Excel workbook."""

from openpyxl.styles import Border, Font, PatternFill, Side

DARK_HEX = "1A1A1A"
MID_HEX = "4A4A4A"
LIGHT_HEX = "888888"
TOTAL_BG = "F2F2F2"  # very light grey for total rows
RULE_HEX = "CCCCCC"


def thin(color: str = RULE_HEX) -> Side:
    return Side(border_style="thin", color=color)


def medium(color: str = DARK_HEX) -> Side:
    return Side(border_style="medium", color=color)


def grand_border() -> Border:
    return Border(top=medium(), bottom=medium())


def body_font(bold: bool = False, size: float = 10, color: str = DARK_HEX) -> Font:
    return Font(name="Calibri", bold=bold, size=size, color=color)


def total_fill() -> PatternFill:
    return PatternFill("solid", fgColor=TOTAL_BG)
